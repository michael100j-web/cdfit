#!/usr/bin/env python3
"""CD Fit for Claude: fit CD and UV melting curves (Tm of CMPs) and plot CD spectra exactly as the CD Fit
Word add-in does, with the same model, fitter, defaults and Prism-style graph.

  python cdfit.py fit DATA [options]   fit a data file (.txt .csv .tsv .dat .xlsx, '-' = stdin), a saved state
                                       (.json) or every CD Fit graph in a Word file (.docx)
  python cdfit.py list FILE.docx       list the CD Fit graphs in a Word file
  python cdfit.py example MODE         print the add-in's example data (melt, uv or spec)

`python cdfit.py fit -h` lists the options.
"""
from __future__ import annotations

import argparse
import json
import math
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import cdfit_engine as E  # noqa: E402

MODE_NAMES = {"melt": "CD melting", "uv": "UV melting", "spec": "CD spectrum"}
DIRECTION_NAMES = {"any": "either direction", "up": "signal rises on unfolding", "down": "signal falls on unfolding"}
_RE_CMP_CONC = re.compile(r"ln\(\s*0\.75\s*\*\s*([0-9]*\.?[0-9]+(?:[eE][-+]?[0-9]+)?)\s*\^\s*2\s*\)", re.I)
_RE_H_LINE = re.compile(r"^[ \t]*H[ \t]*=[ \t]*([-+]?[0-9]*\.?[0-9]+(?:[eE][-+]?[0-9]+)?)[ \t]*$", re.I | re.M)


class UsageError(Exception):
    pass


# ============================ Reading data ============================
def read_text_file(path):
    raw = open(path, "rb").read()
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16")
    for enc in ("utf-8-sig", "cp1250", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", "replace")


def read_xlsx(path, sheet=None):
    try:
        import openpyxl
    except ImportError as e:
        raise UsageError("Reading .xlsx needs openpyxl (pip install openpyxl), or save the sheet as .csv.") from e
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet and sheet not in wb.sheetnames:
        raise UsageError(f"No sheet named {sheet!r}; the sheets are: {', '.join(wb.sheetnames)}")
    ws = wb[sheet] if sheet else wb.worksheets[0]
    lines = []
    for row in ws.iter_rows(values_only=True):
        cells = ["" if v is None else E.js_str(v) if isinstance(v, (int, float)) and not isinstance(v, bool)
                 else str(v).replace("\t", " ").replace("\n", " ") for v in row]
        while cells and cells[-1] == "":
            cells.pop()
        if any(c.strip() for c in cells):
            lines.append("\t".join(cells))
    wb.close()
    return "\n".join(lines)


def read_data_files(paths, sheet, mode):
    """One or more data files as one table: each file's Y columns side by side, matched by X. A JASCO export keeps
    only its first channel (CD), named after the file, as on the CD Fit page."""
    import cdfit_spectra as SP
    parsed, skipped = [], []
    natural = lambda p: [int(t) if t.isdigit() else t.lower() for t in re.split(r"([0-9]+)", os.path.basename(p))]
    for path in sorted(paths, key=natural) if len(paths) > 1 else paths:   # CMP1_4C before CMP1_20C
        if not os.path.exists(path):
            raise UsageError(f"No such file: {path}")
        ext = os.path.splitext(path)[1].lower()
        if ext in (".docx", ".json"):
            raise UsageError(f"{os.path.basename(path)}: a Word file or a saved state can only be fitted on its own.")
        if ext == ".xls":
            raise UsageError("Old .xls files are not supported: save the sheet as .xlsx or .csv.")
        raw = read_xlsx(path, sheet) if ext in (".xlsx", ".xlsm") else read_text_file(path)
        jasco = SP.jasco_text(raw)
        p = E.parse_data(jasco or raw)
        if not p["rows"] or not p["names"]:
            skipped.append(os.path.basename(path))
            continue
        parsed.append((os.path.splitext(os.path.basename(path))[0], p, jasco is not None, raw))
    if skipped:
        print("Skipped (no numbers found): " + ", ".join(skipped), file=sys.stderr)
    if not parsed:
        raise UsageError("No numbers found: the data need an X column and at least one Y column.")
    if len(parsed) == 1 and not parsed[0][2]:
        return parsed[0][3]
    text, _ = SP.merge([(stem, p, jasco) for stem, p, jasco, _ in parsed], mode or "melt")
    return text


def load_sources(paths, sheet=None, mode=None):
    """[(label, state, from_graph)] for data files, a saved state or a Word file."""
    if paths == ["-"]:
        return [("data", {"dataText": sys.stdin.read()}, False)]
    if len(paths) > 1:
        return [("data", {"dataText": read_data_files(paths, sheet, mode)}, False)]
    path = paths[0]
    if not os.path.exists(path):
        raise UsageError(f"No such file: {path}")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        import cdfit_word as W
        graphs = W.read_graphs(path)
        if not graphs:
            raise UsageError("No CD Fit graphs in this Word file (pictures inserted by CD Fit carry its data in "
                             "their alt text).")
        out = []
        for g in graphs:
            if "error" in g:
                print(f"Graph {g['index']}: {g['error']}", file=sys.stderr)
                continue
            out.append((f"graph{g['index']}", g["state"], True))
        return out
    if ext == ".json":
        st = json.loads(read_text_file(path))
        if not isinstance(st, dict) or "dataText" not in st:
            raise UsageError("The .json file is not a CD Fit state (it has no dataText).")
        return [("state", st, True)]
    return [("data", {"dataText": read_data_files([path], sheet, mode)}, False)]


# ============================ Options → state ============================
def set_dh(eq, value):
    if _RE_H_LINE.search(eq):
        return _RE_H_LINE.sub("H=" + E.js_str(value), eq, count=1)
    lines = eq.split("\n")
    at = next((i + 1 for i, l in enumerate(lines) if re.match(r"\s*R\s*=", l, re.I)), None)
    if at is None or not re.search(r"\bH\b", eq, re.I):
        raise UsageError("--dH works with the CMP model, whose ΔH is H.")
    lines.insert(at, "H=" + E.js_str(value))
    return "\n".join(lines)


def free_dh(eq):
    if not _RE_H_LINE.search(eq):
        if re.search(r"\bH\b", eq, re.I):
            return eq            # already fitted
        raise UsageError("--fit-dH works with the CMP model, whose ΔH is H.")
    return _RE_H_LINE.sub("", eq, count=1).replace("\n\n", "\n")


def _kv(text, what):
    if "=" not in text:
        raise UsageError(f"{what} needs NAME=VALUE, got {text!r}")
    k, v = text.split("=", 1)
    return k.strip(), v.strip()


_TRUE = {"1", "true", "yes", "on"}
_FALSE = {"0", "false", "no", "off"}


def apply_options(S, a, from_graph):
    """Command-line options on top of a state, as if they had been set in the add-in's panel."""
    if a.mode and not from_graph:
        S = {**E.base_state(a.mode), "dataText": S.get("dataText", "")}
    S = E.apply_state(S, from_graph)
    if a.preset:
        S["eqText"] = E.PRESETS[a.preset]["text"]
    if a.equation:
        S["eqText"] = read_text_file(a.equation) if os.path.exists(a.equation) else a.equation.replace("\\n", "\n")
    if a.dH is not None and a.fit_dH:
        raise UsageError("Use either --dH or --fit-dH.")
    if a.dH is not None:
        try:
            S["eqText"] = set_dh(S["eqText"], float(a.dH))
        except ValueError as e:
            raise UsageError(f"--dH needs a number in J/mol, got {a.dH!r}") from e
    if a.fit_dH:
        S["eqText"] = free_dh(S["eqText"])
        S["params"].pop("h", None)
    if a.no_fit:
        S["fit"] = False
    if a.force_fit:
        S["fit"] = True
    if a.direction:
        S["stepDir"] = a.direction
    for item in a.start:
        k, v = _kv(item, "--start")
        S["params"][k.lower()] = {"value": v.replace(",", "."), "fixed": False}
    for item in a.fix:
        k, v = _kv(item, "--fix")
        S["params"][k.lower()] = {"value": v.replace(",", "."), "fixed": True}
    for k in a.free:
        S["params"].pop(k.strip().lower(), None)
    if a.ymult is not None:
        S["yMult"] = a.ymult
    if a.xrange:
        S["xFrom"], S["xTo"] = a.xrange
    if a.columns or a.names:
        p = E.parse_data(S.get("dataText", ""))
        E.build_data(S)                      # one settings entry per column, as the add-in keeps them
        if a.columns:
            want = set()
            for tok in a.columns.split(","):
                tok = tok.strip()
                if tok.isdigit() and 1 <= int(tok) <= len(p["names"]):
                    want.add(int(tok) - 1)
                elif tok in p["names"]:
                    want.add(p["names"].index(tok))
                else:
                    raise UsageError(f"--columns: no column {tok!r}; the Y columns are "
                                     + ", ".join(f"{i + 1} {n}" for i, n in enumerate(p["names"])))
            for j, cfg in enumerate(S["series"]):
                cfg["visible"] = j in want
        if a.names:
            names = a.names.split("|") if "|" in a.names else a.names.split(",")
            for j, nm in enumerate(names[:len(S["series"])]):
                if nm.strip():
                    S["series"][j]["name"] = nm.strip()
    fmt = S["fmt"]
    if a.xtitle is not None:
        fmt["xTitle"] = a.xtitle
    if a.ytitle is not None:
        fmt["yTitle"] = a.ytitle
    if a.legend:
        fmt["legend"] = a.legend
    if a.width_cm:
        fmt["exportCm"] = a.width_cm
    if a.dpi:
        fmt["exportFmt"] = "png600" if a.dpi == 600 else "png300"
    for item in a.set:
        k, v = _kv(item, "--set")
        if k not in E.FMT_FIELDS and k not in E.FMT_CHECKS:
            raise UsageError(f"--set: unknown graph setting {k!r}; use one of " + ", ".join(E.FMT_FIELDS + E.FMT_CHECKS))
        if k in E.FMT_CHECKS:
            if v.lower() not in _TRUE | _FALSE:
                raise UsageError(f"--set {k} takes true or false")
            fmt[k] = v.lower() in _TRUE
        else:
            fmt[k] = v
    return S


# ============================ Reporting ============================
def model_description(S, model):
    """One line naming the model and, for the CMP model, whether ΔH is fixed or fitted."""
    eq = S.get("eqText", "")
    preset = next((k for k, p in E.PRESETS.items() if p["text"].strip() == eq.strip()), None)
    if _RE_CMP_CONC.search(eq) and model and "tm" in model.keys:
        h = _RE_H_LINE.search(eq)
        dh = (f"ΔH fixed at {E.fmt_num(float(h.group(1)) / 1000).replace('-', '−')} kJ/mol" if h
              else "ΔH fitted" if "h" in model.keys else "ΔH as defined in the equation")
        name = "Tm of CMPs (trimer ⇌ 3 monomers)" + ("" if preset == "cmp" else ", edited")
        return f"{name}: {dh}"
    if preset:
        return E.PRESETS[preset]["label"]
    return "custom equation" + (f" (parameters {', '.join(model.params)})" if model else "")


def clean(v):
    if isinstance(v, float):
        return None if math.isnan(v) or math.isinf(v) else v
    if isinstance(v, dict):
        return {k: clean(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [clean(x) for x in v]
    return v


def series_report(session, ser):
    S, model = session.S, session.model
    cfg = S["series"][ser["idx"]]
    rec = {"name": ser["name"], "column": ser["idx"] + 1, "points": len(ser["xs"]), "excluded": len(ser["xsEx"])}
    if not cfg.get("visible"):
        rec["status"] = "hidden"
        return rec
    if not S.get("fit"):
        rec["status"] = "plotted"
        rec["bands"] = E.bands(ser)
        return rec
    fit = session.fit_of(ser)
    if not fit:
        rec["status"] = "no points"
        return rec
    if fit.get("error"):
        rec["status"] = "error"
        rec["error"] = fit["error"]
        return rec
    rec["status"] = "fitted"
    rec["params"] = {name: {"value": fit["p"][i], "se": fit["se"][i], "ci95": fit["ci"][i], "fixed": fit["fixed"][i],
                            "text": (E.fmt_num(fit["p"][i]) + " (fixed)" if fit["fixed"][i]
                                     else E.fmt_pm(fit["p"][i], fit["se"][i]))}
                     for i, name in enumerate(model.params)}
    tm = session.tm(ser)
    if tm:
        rec["Tm"] = {"value": tm["v"], "se": tm["se"], "ci95": tm["ci"], "text": E.fmt_pm(tm["v"], tm["se"]) + " °C"}
    rec.update({"r2": fit["r2"], "syx": fit["syx"], "n": fit["n"], "df": fit["df"], "converged": fit["converged"],
                "iterations": fit["iter"], "ambiguous": fit["ambiguous"]})
    if "step" in fit:
        rec.update({"step": fit["step"], "direction": fit.get("dir")})
    if fit.get("avoided"):
        rec["unconstrained"] = fit["avoided"]
    if fit.get("limit"):
        rec["on_constraint"] = True
    return rec


def notes_for(session, ser):
    """The add-in's Results-tab remarks for one fitted data set."""
    fit = session.fit_of(ser)
    if not fit or fit.get("error"):
        return []
    out = []
    if not fit["converged"]:
        out.append(f"not converged in {fit['iter']} iterations")
    if fit["ambiguous"]:
        out.append("ambiguous: the parameters are not identifiable (no standard errors)")
    if fit.get("avoided"):
        av = fit["avoided"]
        what = ("puts Tm outside the data" if av["outside"]
                else f"has a {'falling' if fit.get('dir') == 'up' else 'rising'} step")
        out.append(f"without the direction constraint the fit {what} (Tm = {E.fmt_num(av['tm'], 4)} °C, "
                   f"Sy.x = {E.fmt_num(av['syx'], 3)}); this fit was found by scanning Tm across the data")
    if fit.get("limit"):
        out.append("the best allowed fit lies on the constraint, not at a minimum, so Tm has no standard error; "
                   "the data may not show a transition in this direction")
    return out


def print_report(session, label, files, model_line):
    S = session.S
    head = MODE_NAMES.get(S["mode"], S["mode"])
    if S.get("fit"):
        head += f" · {model_line}"
        if E.step_supported(session.model):
            head += f" · {DIRECTION_NAMES.get(S.get('stepDir'), 'either direction')}"
    else:
        head += " · plotted, no fit"
    print(f"[{label}] {head}" if label else head)
    if session.model_error:
        print(f"  Equation error: {session.model_error}")
    for ser in session.data["series"]:
        cfg = S["series"][ser["idx"]]
        if not cfg.get("visible"):
            continue
        if not S.get("fit"):
            b = E.bands(ser)
            if b:
                pos = f"max {E.fmt_num(b['xMax'], 4)} nm ({E.fmt_num(b['yMax'], 4)}), " if b["hasPos"] else \
                    "no clear positive band, "
                rpn = f", Rpn {E.js_to_fixed(b['rpn'], 3)}" if math.isfinite(b["rpn"]) else ""
                cross = f", crossover {E.js_to_fixed(b['cross'], 1)} nm" if math.isfinite(b["cross"]) else ""
                print(f"  {ser['name']}: {pos}min {E.fmt_num(b['xMin'], 4)} nm ({E.fmt_num(b['yMin'], 4)}){cross}{rpn}")
            continue
        fit = session.fit_of(ser)
        if not fit:
            print(f"  {ser['name']}: no points")
            continue
        if fit.get("error"):
            print(f"  {ser['name']}: fit failed: {fit['error']}")
            continue
        stats = f"R² {E.js_to_fixed(fit['r2'], 4)} · Sy.x {E.fmt_num(fit['syx'], 3)} · N {fit['n']}"
        tm = session.tm(ser)
        if tm:
            ci = tm["ci"]
            ci_txt = (f" (95% CI {E.js_to_fixed(ci[0], 2)} to {E.js_to_fixed(ci[1], 2)})"
                      if ci and math.isfinite(ci[0]) else "")
            fixed = " (fixed)" if tm["fixed"] else ""
            print(f"  {ser['name']}: Tm = {E.fmt_pm(tm['v'], tm['se'])} °C{fixed}{ci_txt} · {stats}")
        else:
            pars = ", ".join(f"{n} = " + (E.fmt_num(fit['p'][i]) + " (fixed)" if fit["fixed"][i]
                                          else E.fmt_pm(fit["p"][i], fit["se"][i]))
                             for i, n in enumerate(session.model.params))
            print(f"  {ser['name']}: {pars} · {stats}")
        for note in notes_for(session, ser):
            print(f"    note: {note}")
    if files:
        print("  files: " + ", ".join(files))


# ============================ Commands ============================
def cmd_fit(a):
    import cdfit_plot as P
    import cdfit_spectra as SP
    import cdfit_word as W

    sources = load_sources(a.sources, a.sheet, a.mode)
    if a.graph:
        pick = {int(t) for t in a.graph.split(",") if t.strip().isdigit()}
        sources = [s for s in sources if s[0].startswith("graph") and int(s[0][5:]) in pick]
        if not sources:
            raise UsageError(f"--graph {a.graph}: no such graph (see: cdfit.py list FILE.docx)")
    first = a.sources[0]
    stem = ("data" if first == "-" else SP.common_stem(a.sources) if len(a.sources) > 1
            else os.path.splitext(os.path.basename(first))[0])
    out_dir = a.out or os.path.join(os.getcwd() if first == "-" else os.path.dirname(os.path.abspath(first)),
                                    stem + "_cdfit")
    os.makedirs(out_dir, exist_ok=True)
    formats = [f.strip().lower().lstrip(".") for f in a.formats.split(",") if f.strip()]
    for f in formats:
        if f not in ("png", "svg", "pdf"):
            raise UsageError(f"--formats: {f!r} is not png, svg or pdf")
    if "png" not in formats and a.docx is not None:
        formats.append("png")
    multi = len(sources) > 1
    report, docx_graphs, status = [], [], 0
    for label, st, from_graph in sources:
        S = apply_options(st, a, from_graph)
        session = E.Session(S)
        if not session.data["series"]:
            raise UsageError("No numbers found: the data need an X column and at least one Y column.")
        base = a.name or stem
        if multi or label.startswith("graph"):
            base = f"{base}_{label}"
        files = []
        paths = [os.path.join(out_dir, f"{base}.{f}") for f in formats]
        W_px, H_px, width_cm = P.render(session, paths)
        files += paths
        rows = session.results_matrix()
        tsv = os.path.join(out_dir, f"{base}_results.tsv")
        with open(tsv, "w", encoding="utf-8-sig", newline="\n") as fh:
            fh.write("\n".join("\t".join(r) for r in rows) + "\n")
        state_path = os.path.join(out_dir, f"{base}_state.json")
        with open(state_path, "w", encoding="utf-8") as fh:
            fh.write(json.dumps(session.S, ensure_ascii=False, indent=1))
        model_line = model_description(session.S, session.model) if session.S.get("fit") else None
        rec = {"source": label, "mode": session.S["mode"], "model": model_line, "equation": session.S.get("eqText"),
               "direction": session.S.get("stepDir") if E.step_supported(session.model) else None,
               "model_error": session.model_error,
               "series": [series_report(session, s) for s in session.data["series"]],
               "notes": {s["name"]: notes_for(session, s) for s in session.visible if notes_for(session, s)},
               "table": rows, "summary": session.summary_text(),
               "files": {"graph": paths, "results_tsv": tsv, "state": state_path}}
        extra = []
        if session.S["mode"] == "spec":   # many spectra: one row each, and θ at the wavelength against temperature
            v = SP.analyse(session, a.wavelength)
            spec_tsv = os.path.join(out_dir, f"{base}_spectra.tsv")
            with open(spec_tsv, "w", encoding="utf-8-sig", newline="\n") as fh:
                fh.write("\n".join("\t".join(r) for r in SP.table(v)) + "\n")
            rec["files"]["spectra_tsv"] = spec_tsv
            rec["spectra"] = [{"name": r["name"], "label": r["label"], "bands": r["bands"], "value_at_wavelength": r["at"]}
                              for r in v["rows"]]
            extra.append(f"spectra table ({len(v['rows'])} spectra, one row each): {spec_tsv}")
            if v["trend"]:
                lam = E.fmt_num(a.wavelength, 4)
                melt_path = os.path.join(out_dir, f"{base}_melt_{lam}nm.txt")
                with open(melt_path, "w", encoding="utf-8") as fh:
                    fh.write(SP.melt_table(v) + "\n")
                rec["files"]["melt_table"] = melt_path
                what = "temperature" if v["celsius"] else "the number in the names"
                # fit it right away with the CMP model, as the page's spectra section does
                unit = re.search(r"\(([^()]*)\)\s*$", session.S["fmt"].get("yTitle") or "")
                melt = E.Session({**E.base_state("melt"), "dataText": SP.melt_table(v),
                                  "fmt": {**E.mode_fmt("melt"), "yTitle": f"*θ*_{{{lam}}}" + (f" ({unit.group(1)})" if unit else ""),
                                          **({} if v["celsius"] else {"xTitle": "Number in the spectrum name"})}})
                melt_png = os.path.join(out_dir, f"{base}_melt_{lam}nm.png")
                P.render(melt, [melt_png])
                rec["files"]["melt_graph"] = melt_png
                rec["spectra_fit"] = [series_report(melt, s) for s in melt.data["series"]]
                parts = []
                for ser in melt.visible:
                    fit, tm = melt.fit_of(ser), melt.tm(ser)
                    parts.append(f"{ser['name']}: fit failed ({fit['error']})" if fit and fit.get("error") else
                                 f"{ser['name']} Tm = {E.fmt_pm(tm['v'], tm['se'])} °C" if tm else ser["name"])
                extra.append(f"θ at {lam} nm against {what}, fitted with the CMP model (ΔH −500 kJ/mol): " + "; ".join(parts))
                extra.append(f"  data: {melt_path}  graph: {melt_png}")
        res_json = os.path.join(out_dir, f"{base}_results.json")
        rec["files"]["results_json"] = res_json
        with open(res_json, "w", encoding="utf-8") as fh:
            fh.write(json.dumps(clean(rec), ensure_ascii=False, indent=1))
        report.append(rec)
        if a.docx is not None:
            png = next(p for p in paths if p.endswith(".png"))
            with open(png, "rb") as fh:
                docx_graphs.append({"png": fh.read(), "width_px": W_px, "height_px": H_px, "width_cm": width_cm,
                                    "alt": session.alt_text(W.state_json(session.S)),
                                    "table": rows, "heading": None, "caption": None})
        if not a.json:
            print_report(session, label if multi or label.startswith("graph") else "",
                         [p for p in paths] + [tsv], model_line)
            for line in extra:
                print("  " + line)
        if session.model_error or (session.S.get("fit") and session.fits
                                   and all(f is None or f.get("error") for f in session.fits)):
            status = 1
    if a.docx is not None:
        docx_path = a.docx or os.path.join(out_dir, f"{a.name or stem}.docx")
        W.write_docx(docx_path, docx_graphs, title=f"CD Fit: {stem}")
        for rec in report:
            rec["files"]["docx"] = docx_path
        if not a.json:
            print(f"Word file (graphs editable in the CD Fit add-in): {docx_path}")
    if a.json:
        print(json.dumps(clean(report if multi else report[0]), ensure_ascii=False, indent=1))
    return status


def cmd_list(a):
    import cdfit_word as W
    graphs = W.read_graphs(a.docx)
    if not graphs:
        print("No CD Fit graphs in this file.")
        return 1
    for g in graphs:
        if "error" in g:
            print(f"{g['index']}. {g['error']}")
            continue
        st = g["state"]
        names = [n for n in E.parse_data(st.get("dataText", ""))["names"]]
        mode = MODE_NAMES.get(st.get("mode", "melt"), st.get("mode"))
        print(f"{g['index']}. {mode} · {len(names)} data set(s) · {g['summary']}")
    return 0


def cmd_example(a):
    text = E.example_data(a.mode)
    if a.out:
        with open(a.out, "w", encoding="utf-8") as fh:
            fh.write(text + "\n")
        print(a.out)
    else:
        print(text)
    return 0


def build_parser():
    p = argparse.ArgumentParser(prog="cdfit.py", description=__doc__.split("\n\n")[0],
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    f = sub.add_parser("fit", help="fit data, a saved state or the CD Fit graphs in a Word file",
                       description="Fit and plot as the CD Fit add-in does. Defaults are the add-in's: CMP model "
                                   "(ΔH −500 kJ/mol), either direction for CD and rising for UV.")
    f.add_argument("sources", nargs="+", metavar="source",
                   help="data (.txt .csv .tsv .dat .xlsx, '-' = stdin; X in the first column, one Y "
                                  "column per sample, optional header row; several files become one table, one column per "
                                  "file), a CD Fit state (.json) or a .docx")
    g = f.add_argument_group("data")
    g.add_argument("--mode", choices=["melt", "uv", "spec"], help="CD melting (default), UV melting or CD spectrum")
    g.add_argument("--sheet", help="sheet of an .xlsx file (default: the first)")
    g.add_argument("--columns", help="Y columns to show and fit: numbers from 1 or header names, comma-separated")
    g.add_argument("--names", help="legend names for the Y columns in order, separated by | or commas")
    g.add_argument("--ymult", help="multiply Y by this, e.g. 0.001 for 10³ deg cm² dmol⁻¹")
    g.add_argument("--xrange", nargs=2, metavar=("MIN", "MAX"),
                   help="fit only points with X inside; the others are drawn faded ('' leaves a side open)")
    g.add_argument("--graph", help="for a .docx: which CD Fit graphs, numbered as in `list` (default: all)")
    g.add_argument("--wavelength", type=float, default=225.0,
                   help="spectra: the wavelength (nm) for the table's value column and for θ against temperature (225)")
    m = f.add_argument_group("model")
    m.add_argument("--preset", choices=list(E.PRESETS), help="cmp (default), twostate (ΔH fitted) or boltz")
    m.add_argument("--equation", help="Prism-syntax equation: a file, or one string with lines separated by \\n")
    m.add_argument("--dH", help="ΔH in the CMP model, fixed, in J/mol, written --dH=-400000 (the preset has -500000)")
    m.add_argument("--fit-dH", action="store_true", help="fit ΔH in the CMP model instead of fixing it")
    m.add_argument("--direction", choices=["any", "up", "down"],
                   help="transition direction: up = signal rises on unfolding (UV), down = falls (CD at 225 nm)")
    m.add_argument("--start", action="append", default=[], metavar="NAME=VALUE", help="starting value")
    m.add_argument("--fix", action="append", default=[], metavar="NAME=VALUE", help="hold a parameter constant")
    m.add_argument("--free", action="append", default=[], metavar="NAME", help="clear a start value or fix")
    fg = m.add_mutually_exclusive_group()
    fg.add_argument("--no-fit", action="store_true", help="plot only")
    fg.add_argument("--fit", dest="force_fit", action="store_true", help="fit even in spectrum mode")
    gr = f.add_argument_group("graph")
    gr.add_argument("--xtitle", help="X axis title (markup: *italic*, ^{sup}, _{sub})")
    gr.add_argument("--ytitle", help="Y axis title")
    gr.add_argument("--legend", choices=["tr", "tl", "br", "bl", "off"])
    gr.add_argument("--width-cm", type=float, help="graph width in the document (default 12)")
    gr.add_argument("--dpi", type=int, choices=[300, 600], help="PNG resolution (default 300)")
    gr.add_argument("--set", action="append", default=[], metavar="KEY=VALUE",
                    help="any add-in graph setting: " + ", ".join(E.FMT_FIELDS + E.FMT_CHECKS))
    o = f.add_argument_group("output")
    o.add_argument("--out", help="output folder (default: <source>_cdfit next to the source)")
    o.add_argument("--name", help="base name of the output files (default: the source's name)")
    o.add_argument("--formats", default="png", help="graph formats, comma-separated: png, svg, pdf (default png)")
    o.add_argument("--docx", nargs="?", const="", metavar="PATH",
                   help="also write a Word file with the graph(s), editable in the CD Fit add-in, and results")
    o.add_argument("--json", action="store_true", help="print the results as JSON instead of text")
    lst = sub.add_parser("list", help="list the CD Fit graphs in a Word file")
    lst.add_argument("docx")
    ex = sub.add_parser("example", help="the add-in's example data")
    ex.add_argument("mode", choices=["melt", "uv", "spec"])
    ex.add_argument("--out", help="write to this file instead of printing")
    return p


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    a = build_parser().parse_args(argv)
    try:
        return {"fit": cmd_fit, "list": cmd_list, "example": cmd_example}[a.cmd](a)
    except UsageError as e:
        print(f"cdfit: {e}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
